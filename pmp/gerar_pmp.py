import sys,datetime as dt,warnings
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import FormulaRule,CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart,LineChart,Reference
from openpyxl.chart.series import SeriesLabel
from openpyxl.comments import Comment
warnings.filterwarnings('ignore')
SRC,OUT,PQ=sys.argv[1],sys.argv[2],sys.argv[3]

# ---------- dados da carteira ----------
df=pd.read_excel(SRC,sheet_name='FGI 8.5-01.02 PLANEJADO',header=7,usecols='B:S').dropna(subset=['OS'])
a=df[(df['PROJETO ATIVO PLANEJAMENTO']=='SIM')&(df['ORIGEM']!='CANCELADO')].copy()
a['m']=pd.to_datetime(a['Mês Planejado'],errors='coerce')
for c in ['CÓDIGO DO PROJETO','SUBESTAÇÃO','TAG']: a[c]=a[c].astype(str).str.strip()
OSL=list(dict.fromkeys(int(x) for x in a['OS']))
cr=pd.read_excel(SRC,sheet_name='Cronograma SET-OUT',header=1,usecols='A:V').dropna(subset=['OS'])
extra={}
for _,r in cr.iterrows():
    alvo=pd.to_datetime(r['Data Alvo'],errors='coerce') if str(r['Meta']).strip().upper()=='EMBALADO' else pd.NaT
    extra[int(r['OS'])]=dict(alvo=None if pd.isna(alvo) else alvo.date(),
        doc='S' if str(r['Data recebimento Documentação']).strip().upper()=='DISPONIVEL' else 'N',
        mat='S' if str(r['Data recebimento Material Critico']).strip().upper()=='RECEBIDO' else 'N')
projs=sorted(a['CÓDIGO DO PROJETO'].unique())

# ---------- estilos ----------
F='Arial'
f_in=Font(name=F); f_n=Font(name=F); f_b=Font(name=F,bold=True)
f_h=Font(name=F,bold=True,color='FFFFFF'); f_t=Font(name=F,bold=True,size=14); f_link=Font(name=F,color='008000')
f_i=Font(name=F,italic=True,size=9)
fill_h=PatternFill('solid',fgColor='1F3864'); fill_in=PatternFill('solid',fgColor='DDEBF7')
fill_sub=PatternFill('solid',fgColor='D9E1F2'); fill_tot=PatternFill('solid',fgColor='EDEDED')
STF={'MEC':'DDEBF7','ELE':'FCE4D6','PLAT':'E2EFDA','CDP':'FFF2CC','NORM':'EDE2F6'}
thin=Side(style='thin',color='BFBFBF'); box=Border(left=thin,right=thin,top=thin,bottom=thin)
center=Alignment(horizontal='center',vertical='center',wrap_text=True)
MFMT='mmm/yy;;'   # mês: zero aparece em branco
HFMT='0;-0;-'
NM=9; START=dt.date(2026,10,1)
MES=[dt.date(START.year+(START.month-1+i)//12,(START.month-1+i)%12+1,1) for i in range(NM)]
def hdr(ws,r,c,v,fill=fill_h,font=f_h):
    x=ws.cell(r,c,v); x.font=font; x.fill=fill; x.alignment=center; x.border=box; return x
def setw(ws,d):
    for k,v in d.items(): ws.column_dimensions[k].width=v
wb=Workbook()

# ---------- LEIAME ----------
R=wb.active; R.title='LEIAME'

# ---------- CARTEIRA (destino do Power Query) ----------
K=wb.create_sheet('CARTEIRA')
kh=['CHAVE','OS','CÓDIGO DO PROJETO','SUBESTAÇÃO','TAG','TIPO PNL','COMPLEXIDADE','ORIGEM','ETAPA','Mês Planejado','Horas / Mês','GIGA']
for j,h in enumerate(kh): hdr(K,1,1+j,h)
for i,(_,r) in enumerate(a.iterrows()):
    rr=2+i
    K.cell(rr,1,f'=B{rr}&I{rr}')
    vals=[int(r['OS']),r['CÓDIGO DO PROJETO'],r['SUBESTAÇÃO'],r['TAG'],r['TIPO PNL'],r['COMPLEXIDADE'],r['ORIGEM'],r['ETAPA'],
          None if pd.isna(r['m']) else r['m'].date(), float(r['Horas / Mês']), r['GIGA']]
    for j,v in enumerate(vals):
        x=K.cell(rr,2+j,v)
        if j==8: x.number_format='mmm/yy'
KLAST=1+len(a)
for c in range(1,13):
    for rr in range(2,KLAST+1): K.cell(rr,c).font=f_n
setw(K,{'A':18,'B':8,'C':40,'D':24,'E':16,'F':9,'G':12,'H':9,'I':15,'J':10,'K':10,'L':7})
K.freeze_panes='A2'; K.auto_filter.ref=f'A1:L{KLAST}'

# ---------- PARAMETROS ----------
P=wb.create_sheet('PARAMETROS')
P['A1']='PARÂMETROS DE CAPACIDADE'; P['A1'].font=f_t
P['A2']='Fundo azul claro = entrada. Mesma conta da aba CAPACIDADE (2) da carteira: dias × horas/dia × eficiência × pessoas.'; P['A2'].font=f_i
hdr(P,4,1,'Item'); hdr(P,4,2,'Unid.')
for i in range(NM):
    c=3+i; x=P.cell(4,c,MES[0] if i==0 else f'=EDATE({L(c-1)}4,1)')
    x.number_format='mmm/yy'; x.font=f_in if i==0 else f_h; x.fill=fill_in if i==0 else fill_h; x.alignment=center; x.border=box
P['C4'].comment=Comment('Mês inicial do PMP (o mês atual). Os outros seguem sozinhos.','PMP')
param=[('Horas por dia','h',[9]*9),('Eficiência','%',[0.78,0.78,0.78,0.9,0.9,0.9,0.9,0.9,0.9]),('Dias úteis','dias',[21,19,15,15,18,22,21,20,22]),
 ('Mecânica','pessoas',[3]*9),('Elétrica (próprios)','pessoas',[2]*9),('Elétrica (terceiros)','pessoas',[9]*9),('Plataforma (giga)','pessoas',[2]*9),
 ('CDP (testes)','pessoas',[3]*9),('Normalização/Insp./Emb.','pessoas',[2]*9)]
for k,(name,u,vals) in enumerate(param):
    r=5+k; P.cell(r,1,name).font=f_n; P.cell(r,2,u).font=f_n
    for i,v in enumerate(vals):
        x=P.cell(r,3+i,v); x.font=f_in; x.fill=fill_in; x.border=box; x.number_format='0%' if u=='%' else '0.0' if u=='pessoas' else '0'
for c in range(6,12):
    P.cell(7,c).comment=Comment('Premissa minha: jan = 15 dias (férias coletivas como jan/26); demais = calendário menos feriados nacionais. Ajuste.','PMP')
    for rr in range(8,14): P.cell(rr,c).comment=Comment('Premissa minha: a carteira não tinha headcount de 2027 — repeti dez/26. Ajuste.','PMP')
P['A15']='Fator capacidade flex (hora extra)'; P['A15'].font=f_n
P['C15']=1.2; P['C15'].font=f_in; P['C15'].fill=fill_in; P['C15'].number_format='0.00'
hdr(P,17,1,'CAPACIDADE (h/mês)'); hdr(P,17,2,'Centro')
for i in range(NM):
    x=P.cell(17,3+i,f'={L(3+i)}4'); x.number_format='mmm/yy'; x.font=f_h; x.fill=fill_h; x.alignment=center
caps=[('Mecânica','MEC','{c}8'),('Elétrica','ELE','({c}9+{c}10)'),('Plataforma','PLAT','{c}11'),('CDP','CDP','{c}12'),('Normalização','NORM','{c}13')]
CAPROW={}
for k,(name,code,ref) in enumerate(caps):
    rr=18+k; CAPROW[code]=rr; P.cell(rr,1,name).font=f_n; P.cell(rr,2,code).font=f_b
    for i in range(NM):
        c=L(3+i); x=P.cell(rr,3+i,f'={c}5*{c}6*{c}7*'+ref.format(c=c)); x.number_format='#,##0'; x.font=f_n; x.border=box
P.cell(23,1,'TOTAL').font=f_b
for i in range(NM):
    c=L(3+i); x=P.cell(23,3+i,f'=SUM({c}18:{c}22)'); x.number_format='#,##0'; x.font=f_b; x.fill=fill_tot; x.border=box
setw(P,{'A':36,'B':10,**{L(3+i):10 for i in range(NM)}}); P.freeze_panes='C5'

# ---------- PMP ----------
M=wb.create_sheet('PMP',1)
M['A1']='PMP — PLANO MESTRE DE PRODUÇÃO'; M['A1'].font=f_t
M['A2']='Você digita só nas células de fundo AZUL CLARO: a OS (coluna A), as premissas (data alvo, doc, material, início da montagem) e, em cada etapa, as horas REALIZADAS e se CONCLUIU (S). O resto vem da aba CARTEIRA por PROCV.'; M['A2'].font=f_i
base=['OS','Projeto','Subestação','TAG','Origem','Giga','Tipo','Complexidade','Data alvo (entrega)','Doc. liberada (S/N)','Material crítico OK (S/N)','Início montagem']
CALVO,CDOC,CMAT,CINI,CGIGA=L(9),L(10),L(11),L(12),L(6)
BASEN=len(base)
for j,h in enumerate(base): hdr(M,5,1+j,h)
hdr(M,4,1,'PAINEL (vem da carteira)'); M.merge_cells(start_row=4,start_column=1,end_row=4,end_column=8)
hdr(M,4,9,'PREMISSAS'); M.merge_cells(start_row=4,start_column=9,end_row=4,end_column=12)
STG=['MEC','ELE','CDP','PLAT','NORM']; ETK={'MEC':'MEC','ELE':'ELE','PLAT':'PLAT','CDP':'CDP','NORM':'NORM/INSP/EMB'}
SUB=['Mês','H. plan.','H. real.','Concl. (S/N)','Saldo']
SC={}; c=BASEN+1
for s in STG:
    SC[s]=dict(zip(['mes','plan','real','concl','saldo'],range(c,c+5)))
    M.merge_cells(start_row=4,start_column=c,end_row=4,end_column=c+4)
    hdr(M,4,c,s if s!='NORM' else 'NORM/INSP/EMB',fill=PatternFill('solid',fgColor=STF[s]),font=f_b)
    for k,h in enumerate(SUB): hdr(M,5,c+k,h,fill=PatternFill('solid',fgColor=STF[s]),font=f_b)
    c+=5
calc=['H. plan. total','H. real. total','Saldo total','Mês de saída','Ordem das etapas','Giga x Plataforma','Início vs programado','Entrega no prazo?','Pronto p/ iniciar?','Ação PCP']
CC={h:c+k for k,h in enumerate(calc)}
for h,cc in CC.items(): hdr(M,5,cc,h)
hdr(M,4,c,'TOTAIS E VERIFICAÇÕES'); M.merge_cells(start_row=4,start_column=c,end_row=4,end_column=c+len(calc)-1)
LASTCOL=c+len(calc)-1
ROW0=6; MAXR=800
yn=DataValidation(type='list',formula1='"S,N"',allow_blank=True); M.add_data_validation(yn)
KR='CARTEIRA!$A:$L'; KB='CARTEIRA!$B:$L'
for i in range(MAXR-ROW0+1):
    rr=ROW0+i; os_=OSL[i] if i<len(OSL) else None
    x=M.cell(rr,1,os_); x.font=f_in
    for j,col in enumerate([2,3,4,7,5,6]):   # Projeto,SE,TAG,Origem,Tipo,Compl -> col idx em CARTEIRA!B:K
        tgt=[2,3,4,5,6,7][j]
    # Projeto(2) SE(3) TAG(4) Origem(5) Tipo(6) Compl(7)  -> índices em B:K: proj 2, se 3, tag 4, tipo 5, compl 6, origem 7
    for colM,idx in [(2,2),(3,3),(4,4),(5,7),(6,11),(7,5),(8,6)]:
        M.cell(rr,colM,f'=IF($A{rr}="","",IFERROR(VLOOKUP($A{rr},{KB},{idx},FALSE),""))').font=f_n
    e=extra.get(os_,{}) if os_ else {}
    x=M.cell(rr,9,e.get('alvo')); x.font=f_in; x.number_format='dd/mm/yy'
    M.cell(rr,10,e.get('doc')).font=f_in; M.cell(rr,11,e.get('mat')).font=f_in
    yn.add(M.cell(rr,10)); yn.add(M.cell(rr,11))
    x=M.cell(rr,12); x.font=f_in; x.number_format='dd/mm/yy'
    for s in STG:
        d=SC[s]; key=f'$A{rr}&"{ETK[s]}"'
        x=M.cell(rr,d['mes'],f'=IF($A{rr}="",0,IFERROR(VLOOKUP({key},{KR},10,FALSE),0))'); x.number_format=MFMT; x.font=f_n
        if s=='PLAT':
            x=M.cell(rr,d['plan'],f'=IF(${CGIGA}{rr}="SIM",IFERROR(VLOOKUP({key},{KR},11,FALSE),0),0)')
        else:
            x=M.cell(rr,d['plan'],f'=IF($A{rr}="",0,IFERROR(VLOOKUP({key},{KR},11,FALSE),0))'); x.number_format=HFMT; x.font=f_n
        x=M.cell(rr,d['real']); x.font=f_in; x.number_format='0'
        x=M.cell(rr,d['concl']); x.font=f_in; x.alignment=Alignment(horizontal='center'); yn.add(x)
        P_,R_,C_=L(d['plan']),L(d['real']),L(d['concl'])
        x=M.cell(rr,d['saldo'],f'=IF({C_}{rr}="S",0,MAX(0,{P_}{rr}-{R_}{rr}))'); x.number_format=HFMT; x.font=f_b
    def sm(k): return '+'.join(f'{L(SC[s][k])}{rr}' for s in STG)
    M.cell(rr,CC['H. plan. total'],f'={sm("plan")}').number_format=HFMT
    M.cell(rr,CC['H. real. total'],f'={sm("real")}').number_format=HFMT
    M.cell(rr,CC['Saldo total'],f'={sm("saldo")}').number_format=HFMT
    ms=','.join(f'{L(SC[s]["mes"])}{rr}' for s in STG)
    M.cell(rr,CC['Mês de saída'],f'=MAX({ms})').number_format=MFMT
    m={s:f'{L(SC[s]["mes"])}{rr}' for s in STG}
    pairs=[(STG[x1],STG[x2]) for x1 in range(5) for x2 in range(x1+1,5)]
    conds=','.join(f'AND({m[x1]}>0,{m[x2]}>0,{m[x1]}>{m[x2]})' for x1,x2 in pairs)
    M.cell(rr,CC['Ordem das etapas'],f'=IF($A{rr}="","",IF(OR({conds}),"ERRO DE ORDEM","OK"))')
    sa=f'{L(CC["Mês de saída"])}{rr}'
    M.cell(rr,CC['Entrega no prazo?'],f'=IF(OR(${CALVO}{rr}="",{sa}=0),"",IF({sa}>DATE(YEAR(${CALVO}{rr}),MONTH(${CALVO}{rr}),1),"ATRASA","NO PRAZO"))')
    pk=f'IFERROR(VLOOKUP($A{rr}&"PLAT",{KR},11,FALSE),0)'
    M.cell(rr,CC['Giga x Plataforma'],f'=IF($A{rr}="","",IF(AND(${CGIGA}{rr}="SIM",{pk}=0),"FALTA PLAT",IF(AND(${CGIGA}{rr}<>"SIM",{pk}>0),"ERRO: PLAT SEM GIGA","OK")))')
    mm=m['MEC']; ini=f'${CINI}{rr}'
    M.cell(rr,CC['Início vs programado'],f'=IF(OR($A{rr}="",{mm}=0),"",IF({ini}="",IF(EOMONTH({mm},0)<TODAY(),"NÃO INICIOU",""),IF({ini}>EOMONTH({mm},0),"INICIOU ATRASADO",IF({ini}<{mm},"ADIANTADO","NO PRAZO"))))')
    mec=m['MEC']; smec=f'{L(SC["MEC"]["saldo"])}{rr}'
    M.cell(rr,CC['Pronto p/ iniciar?'],f'=IF($A{rr}="","",IF(AND(${CDOC}{rr}="S",${CMAT}{rr}="S"),"LIBERADO",IF({smec}=0,"-",IF(AND({mec}>0,{mec}<=EDATE(PARAMETROS!$C$4,1)),"RISCO","PENDENTE"))))')
    past=','.join(f'AND({L(SC[s]["mes"])}{rr}>0,{L(SC[s]["mes"])}{rr}<PARAMETROS!$C$4,{L(SC[s]["saldo"])}{rr}>0)' for s in STG)
    nom=','.join(f'AND({L(SC[s]["mes"])}{rr}=0,{L(SC[s]["saldo"])}{rr}>0)' for s in STG)
    M.cell(rr,CC['Ação PCP'],f'=IF($A{rr}="","",IF(OR({past}),"REPROGRAMAR",IF(OR({nom}),"SEM MÊS","")))')
    for cc in range(1,LASTCOL+1):
        x=M.cell(rr,cc); x.border=box
        if cc>=CC['H. plan. total']: x.font=f_n
    for s in STG:
        M.cell(rr,SC[s]['real']).fill=fill_in; M.cell(rr,SC[s]['concl']).fill=fill_in
    for cc in (1,9,10,11,12): M.cell(rr,cc).fill=fill_in
NR=len(OSL)
red=PatternFill('solid',fgColor='F8CBAD'); amb=PatternFill('solid',fgColor='FFE699'); grn=PatternFill('solid',fgColor='C6EFCE')
rng=f'{L(CC["Ordem das etapas"])}{ROW0}:{L(LASTCOL)}{MAXR}'
for txt,fl in [('FALTA PLAT',red),('ERRO: PLAT SEM GIGA',red),('INICIOU ATRASADO',red),('NÃO INICIOU',red),('ADIANTADO',grn),('ERRO DE ORDEM',red),('ATRASA',red),('RISCO',red),('REPROGRAMAR',amb),('SEM MÊS',amb),('PENDENTE',amb),('LIBERADO',grn),('NO PRAZO',grn)]:
    M.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=[f'"{txt}"'],fill=fl))
# notas explicativas nos cabeçalhos
notes={'Ordem das etapas':'Confere se os meses respeitam a sequência da fábrica: MEC → ELE → CDP → PLAT → NORM/INSP/EMB. "ERRO DE ORDEM" = alguma etapa foi programada para um mês ANTES da etapa que vem antes dela (ex.: ELE em out e MEC em nov).',
 'Giga x Plataforma':'Giga = SIM → tem que ter horas de Plataforma na carteira ("FALTA PLAT" se não tiver). Giga = NÃO → Plataforma tem que ser 0 ("ERRO: PLAT SEM GIGA" se a carteira tiver horas; essas horas NÃO entram na carga).',
 'Início vs programado':'Compara o Início montagem (real) com o mês programado da MEC. NO PRAZO = começou dentro do mês. INICIOU ATRASADO = começou depois do mês. ADIANTADO = antes. NÃO INICIOU = o mês da MEC já acabou e não tem data de início.',
 'Entrega no prazo?':'Compara o mês da última etapa (Mês de saída) com a Data alvo. "ATRASA" = pelo plano o painel termina depois do mês da data alvo. Em branco = sem data alvo.',
 'Pronto p/ iniciar?':'LIBERADO = Doc e Material = S. RISCO = falta doc ou material e a MEC está programada para este mês ou o próximo. PENDENTE = falta doc/material mas a MEC ainda está longe. "-" = MEC já concluída/sem saldo.',
 'Ação PCP':'REPROGRAMAR = alguma etapa tem saldo de horas num mês que já passou (não foi feita no mês planejado). SEM MÊS = etapa tem horas mas não tem mês na carteira. Vazio = nada a fazer.',
 'Saldo total':'Horas que ainda faltam fazer neste painel (é o que conta como carga na aba CARGA x CAPACIDADE).',
 'Mês de saída':'Mês da última etapa programada = mês em que o painel fica pronto para expedir.'}
for h,t in notes.items(): M.cell(5,CC[h]).comment=Comment(t,'PMP',width=320,height=150)
for s in STG:
    M.cell(5,SC[s]['saldo']).comment=Comment('Saldo = H. plan. − H. real. (nunca negativo). Se Concl. = S, saldo = 0.','PMP')
    M.cell(5,SC[s]['real']).comment=Comment('Digite as horas realmente gastas nesta etapa (acumulado).','PMP')
M.freeze_panes=M.cell(ROW0,5)
M.auto_filter.ref=f'A5:{L(LASTCOL)}{ROW0+NR-1}'
setw(M,{'A':7,'B':34,'C':20,'D':14,'E':8,'F':6,'G':7,'H':10,'I':10,'J':8,'K':9,'L':10})
for s in STG:
    for k,w in zip(['mes','plan','real','concl','saldo'],(8,6,6,6,6)): M.column_dimensions[L(SC[s][k])].width=w
for h,w in zip(calc,(7,7,7,8,11,12,12,10,10,12)): M.column_dimensions[L(CC[h])].width=w
M.column_dimensions['G'].hidden=True; M.column_dimensions['H'].hidden=True
M.row_dimensions[5].height=42

# ---------- CARGA x CAPACIDADE ----------
C=wb.create_sheet('CARGA x CAPACIDADE',2)
C['A1']='CARGA × CAPACIDADE POR CENTRO (horas)'; C['A1'].font=f_t
C['A2']='Carga = SALDO de horas do PMP (planejado − realizado) no mês. Backlog acumulado = o que não coube e escorrega pro mês seguinte. Verde ≤85% · Amarelo ≤100% · Laranja ≤ flex · Vermelho > flex.'; C['A2'].font=f_i
hdr(C,4,1,'Centro'); hdr(C,4,2,'Linha'); hdr(C,4,3,'Sem mês'); hdr(C,4,4,'Meses passados')
for i in range(NM):
    x=C.cell(4,5+i,f'=PARAMETROS!{L(3+i)}4'); x.number_format='mmm/yy'; x.font=f_h; x.fill=fill_h; x.alignment=center
LC=5+NM; hdr(C,4,LC,'Total horizonte')
names={'MEC':'Mecânica','ELE':'Elétrica','PLAT':'Plataforma','CDP':'CDP','NORM':'Normalização'}
R1,R2=ROW0,MAXR
def rngc(col): return f'PMP!${L(col)}${R1}:${L(col)}${R2}'
r=5; UT=[]; BLK={}
for s in STG+['TOTAL']:
    BLK[s]=r
    labels=['Capacidade','Capacidade flex','Carga (saldo)','Utilização','Folga (cap − carga)','Backlog acumulado']
    for k,lab in enumerate(labels):
        rr=r+k
        C.cell(rr,1,names.get(s,'TOTAL') if k==0 else None).font=f_b
        C.cell(rr,2,lab).font=f_b if k in (2,3) else f_n
        for i in range(NM):
            col=5+i; cl=L(col)
            if s!='TOTAL':
                d=SC[s]
                fm={0:f'=PARAMETROS!{L(3+i)}{CAPROW[s]}',1:f'={cl}{r}*PARAMETROS!$C$15',
                    2:f'=SUMIFS({rngc(d["saldo"])},{rngc(d["mes"])},{cl}$4)'}
            else:
                fm={k2:'='+'+'.join(f'{cl}{BLK[x]+k2}' for x in STG) for k2 in (0,1,2)}
            fm[3]=f'=IFERROR({cl}{r+2}/{cl}{r},0)'; fm[4]=f'={cl}{r}-{cl}{r+2}'
            fm[5]=f'=MAX(0,$D{r+2}+{cl}{r+2}-{cl}{r})' if i==0 else f'=MAX(0,{L(col-1)}{rr}+{cl}{r+2}-{cl}{r})'
            x=C.cell(rr,col,fm[k]); x.number_format='0%' if k==3 else '#,##0;(#,##0);-'; x.font=f_link if k==0 else f_n; x.border=box
        if k==2:
            if s!='TOTAL':
                d=SC[s]
                C.cell(rr,3,f'=SUMIFS({rngc(d["saldo"])},{rngc(d["mes"])},0)')
                C.cell(rr,4,f'=SUMIFS({rngc(d["saldo"])},{rngc(d["mes"])},">0",{rngc(d["mes"])},"<"&PARAMETROS!$C$4)')
            else:
                C.cell(rr,3,'='+'+'.join(f'C{BLK[x]+2}' for x in STG)); C.cell(rr,4,'='+'+'.join(f'D{BLK[x]+2}' for x in STG))
            for cc in (3,4): C.cell(rr,cc).number_format='#,##0;(#,##0);-'; C.cell(rr,cc).border=box; C.cell(rr,cc).font=f_n
            for cc in range(1,LC+1): C.cell(rr,cc).fill=fill_sub
        if k in (0,1,2,4):
            x=C.cell(rr,LC,f'=SUM(E{rr}:{L(4+NM)}{rr})'); x.number_format='#,##0;(#,##0);-'; x.font=f_b; x.border=box
        if k==3:
            x=C.cell(rr,LC,f'=IFERROR({L(LC)}{r+2}/{L(LC)}{r},0)'); x.number_format='0%'; x.font=f_b; x.border=box; UT.append(rr)
    r+=7
for rr in UT:
    ref=f'E{rr}:{L(LC)}{rr}'
    for f,col in [(f'E{rr}>PARAMETROS!$C$15','FF7C80'),(f'E{rr}>1','F4B183'),(f'E{rr}>0.85','FFE699'),(f'E{rr}>0','C6EFCE')]:
        C.conditional_formatting.add(ref,FormulaRule(formula=[f],fill=PatternFill('solid',fgColor=col),stopIfTrue=True))
# conferência carteira x PMP
r+=0
hdr(C,r,1,'CONFERÊNCIA'); C.merge_cells(start_row=r,start_column=1,end_row=r,end_column=4)
C.cell(r+1,1,'Horas na aba CARTEIRA').font=f_n; CF0=r; C.cell(r+1,4,'=SUM(CARTEIRA!K:K)').number_format='#,##0'
C.cell(r+2,1,'Horas planejadas no PMP').font=f_n; C.cell(r+2,4,f'=SUM({rngc(CC["H. plan. total"])})').number_format='#,##0'
C.cell(r+3,1,'Horas de PLAT ignoradas (Giga = NÃO)').font=f_n
C.cell(r+3,4,'=SUMIFS(CARTEIRA!K:K,CARTEIRA!I:I,"PLAT",CARTEIRA!L:L,"<>SIM")').number_format='#,##0;(#,##0);-'
r+=1
C.cell(r+3,1,'Diferença (≠ 0 → tem OS na carteira que não está no PMP, ou OS+etapa duplicada)').font=f_b
x=C.cell(r+3,4,f'=D{r}-D{r+1}-D{r+2}'); x.number_format='#,##0;(#,##0);-'; x.font=f_b
C.conditional_formatting.add(f'D{r+3}',CellIsRule(operator='notEqual',formula=['0'],fill=red))
CONF=r
r+=5
# aderência ao tempo padrão
hdr(C,r,1,'ADERÊNCIA AO TEMPO PADRÃO'); C.merge_cells(start_row=r,start_column=1,end_row=r,end_column=2)
for j,h in enumerate(['H. plan. (etapas concl.)','H. real. (etapas concl.)','Real ÷ Plan.']): hdr(C,r,3+j,h)
C.cell(r,3).comment=Comment('Só conta etapas com Concl. = S. Real ÷ Plan. acima de 100% = a etapa leva mais tempo que o padrão da aba TEMPOS da carteira → revisar o tempo padrão ou a produtividade.','PMP',width=300,height=120)
for k,s in enumerate(STG):
    rr=r+1+k; d=SC[s]
    C.cell(rr,1,names[s]).font=f_n
    C.cell(rr,3,f'=SUMIFS({rngc(d["plan"])},{rngc(d["concl"])},"S")').number_format='#,##0;(#,##0);-'
    C.cell(rr,4,f'=SUMIFS({rngc(d["real"])},{rngc(d["concl"])},"S")').number_format='#,##0;(#,##0);-'
    x=C.cell(rr,5,f'=IFERROR(D{rr}/C{rr},"")'); x.number_format='0%'
setw(C,{'A':14,'B':20,'C':12,'D':12,**{L(5+i):10 for i in range(NM)},L(LC):13})
C.freeze_panes='E5'
cr0=r+8
for n,s in enumerate(STG+['TOTAL']):
    b=BLK[s]
    bar=BarChart(); bar.type='col'; bar.title=f'{names.get(s,"TOTAL")} — carga × capacidade'
    bar.add_data(Reference(C,min_col=4,max_col=4+NM,min_row=b+2),titles_from_data=False,from_rows=True)
    bar.series[0].tx=SeriesLabel(v='Carga (inclui meses passados)')
    bar.set_categories(Reference(C,min_col=4,max_col=4+NM,min_row=4))
    ln=LineChart()
    for k,lab in ((0,'Capacidade'),(1,'Capacidade flex')):
        ln.add_data(Reference(C,min_col=4,max_col=4+NM,min_row=b+k),titles_from_data=False,from_rows=True); ln.series[-1].tx=SeriesLabel(v=lab)
    bar+=ln; bar.height=7; bar.width=16; bar.legend.position='b'
    C.add_chart(bar,f'{"A" if n%2==0 else "H"}{cr0+(n//2)*15}')

# ---------- PMP MENSAL ----------
S=wb.create_sheet('PMP MENSAL',3)
S['A1']='PMP MENSAL — painéis prontos para expedir por projeto e mês'; S['A1'].font=f_t
S['A2']='Usa a coluna "Mês de saída" do PMP. Projeto novo: acrescente o nome na coluna A (igual ao da carteira) e copie a linha de cima.'; S['A2'].font=f_i
sa=CC['Mês de saída']
def block(top,title,kind):
    hdr(S,top,1,title); hdr(S,top,2,'Meses passados')
    for i in range(NM):
        x=S.cell(top,3+i,f'=PARAMETROS!{L(3+i)}4'); x.number_format='mmm/yy'; x.font=f_h; x.fill=fill_h; x.alignment=center
    hdr(S,top,3+NM,'Total')
    for j,p in enumerate(projs):
        rr=top+1+j; S.cell(rr,1,p).font=f_in; S.cell(rr,1).border=box
        for i in range(-1,NM):
            col=3+i; cl=L(col)
            crit=[f'">0"',f'"<"&PARAMETROS!$C$4'] if i<0 else [f'{cl}${top}']
            if kind=='cnt':
                f='=COUNTIFS('+f'{rngc(2)},$A{rr},'+','.join(f'{rngc(sa)},{c}' for c in crit)+')'
            else:
                f='='+'+'.join('SUMIFS('+f'{rngc(SC[s]["saldo"])},{rngc(2)},$A{rr},'+','.join(f'{rngc(SC[s]["mes"])},{c}' for c in crit)+')' for s in STG)
            x=S.cell(rr,col,f); x.number_format='#,##0;(#,##0);-'; x.font=f_n; x.border=box
        x=S.cell(rr,3+NM,f'=SUM(B{rr}:{L(2+NM)}{rr})'); x.font=f_b; x.number_format='#,##0;(#,##0);-'; x.border=box
    tr=top+1+len(projs); S.cell(tr,1,'TOTAL').font=f_b
    for col in range(2,4+NM):
        x=S.cell(tr,col,f'=SUM({L(col)}{top+1}:{L(col)}{tr-1})'); x.font=f_b; x.fill=fill_tot; x.number_format='#,##0;(#,##0);-'; x.border=box
    return tr
t1=block(4,'Projeto — PAINÉIS prontos','cnt')
block(t1+3,'Projeto — SALDO de horas','h')
setw(S,{'A':52,'B':11,**{L(3+i):9 for i in range(NM)},L(3+NM):9}); S.freeze_panes='B5'

# ---------- LEIAME ----------
pq=open(PQ,encoding='utf-8').read()
txt=[('COMO FUNCIONA ESTE PMP',f_t),('',f_n),
('FLUXO: Carteira (aba PLANEJADO) → aba CARTEIRA (Power Query ou colar) → aba PMP (PROCV) → CARGA x CAPACIDADE e PMP MENSAL.',f_b),
('A carteira continua sendo onde você PLANEJA (muda o mês de uma etapa lá). O PMP só lê a carteira e soma o que você REALIZOU.',f_n),('',f_n),
('ABA PMP — o que cada coluna faz',f_b),
('• OS (fundo azul claro): a única coisa que você digita para um painel novo. Projeto, Subestação, TAG, Origem, Tipo e Complexidade vêm da CARTEIRA por PROCV.',f_n),
('• Data alvo / Doc. liberada / Material crítico / Início montagem: suas premissas (fundo azul claro). Giga vem da carteira.',f_n),
('• Em cada etapa: Mês e H. plan. vêm da carteira (PROCV pela chave OS&ETAPA). H. real. e Concl. (S/N) você digita. Saldo = plan − real (zera se Concl. = S).',f_n),
('• Ordem das etapas: ERRO DE ORDEM se uma etapa foi programada antes da anterior (MEC → ELE → CDP → PLAT → NORM/INSP/EMB).',f_n),
('• Giga x Plataforma: Giga SIM exige horas de PLAT (FALTA PLAT); Giga NÃO exige PLAT = 0 (ERRO: PLAT SEM GIGA — essas horas não entram na carga).',f_n),
('• Início vs programado: compara a data real de início da montagem com o mês programado da MEC (NO PRAZO / INICIOU ATRASADO / ADIANTADO / NÃO INICIOU).',f_n),
('• Entrega no prazo?: ATRASA se o mês de saída (última etapa) é depois do mês da data alvo.',f_n),
('• Pronto p/ iniciar?: LIBERADO (doc e material S) · RISCO (falta algo e a MEC é este mês ou o próximo) · PENDENTE (falta algo, mas há tempo).',f_n),
('• Ação PCP: REPROGRAMAR (sobrou saldo em mês que já passou) · SEM MÊS (etapa com horas mas sem mês na carteira).',f_n),('',f_n),
('COMO LIGAR NA CARTEIRA (Power Query) — faz uma vez',f_b),
('1. Dados > Obter Dados > De Outras Fontes > Consulta Nula.',f_n),
('2. No Editor do Power Query: Página Inicial > Editor Avançado. Apague tudo e cole o código abaixo. Troque o Caminho pelo caminho da sua carteira. Concluído.',f_n),
('3. Renomeie a consulta para CARTEIRA. Página Inicial > Fechar e Carregar Para… > Tabela > Planilha Existente: =CARTEIRA!$A$1 (apague antes os dados que estão lá).',f_n),
('4. Daí pra frente: Dados > Atualizar Tudo. O PMP se atualiza sozinho.',f_n),
('   Carteira no SharePoint/OneDrive? Abra a carteira no Excel desktop, Arquivo > Informações > Copiar Caminho, e use Web.Contents("<link>") no lugar de File.Contents(Caminho).',f_n),
('   Sem Power Query: copie as colunas da aba PLANEJADO e cole como valores na aba CARTEIRA, na mesma ordem (B:L). A coluna A (CHAVE) é =B2&I2.',f_n),
('',f_n),('CÓDIGO POWER QUERY (M):',f_b)]
txt+= [(l,Font(name='Consolas',size=9)) for l in pq.splitlines()]
txt+=[('',f_n),('CONFERÊNCIA',f_b),('Na aba CARGA x CAPACIDADE, o bloco CONFERÊNCIA compara as horas da CARTEIRA com as do PMP. Se a diferença ≠ 0, há OS na carteira que ainda não está no PMP (adicione a OS na coluna A).',f_n),
('Premissas minhas (comentários na aba PARAMETROS): headcount e dias úteis de 2027.',f_n)]
for i,(t,fn) in enumerate(txt): R.cell(1+i,1,t).font=fn
R.column_dimensions['A'].width=150
wb.active=1
wb.save(OUT); print('ok',NR,len(a))

import json,datetime as dt,sys
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter as L
from openpyxl.formatting.rule import FormulaRule,CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.chart import BarChart,LineChart,Reference
from openpyxl.comments import Comment
D=json.load(open('pmp_data.json'))
OUT=sys.argv[1]
wb=Workbook()
F='Arial'
f_in=Font(name=F,color='0000FF'); f_n=Font(name=F); f_b=Font(name=F,bold=True)
f_h=Font(name=F,bold=True,color='FFFFFF'); f_t=Font(name=F,bold=True,size=14)
f_link=Font(name=F,color='008000')
fill_h=PatternFill('solid',fgColor='1F3864'); fill_in=PatternFill('solid',fgColor='FFF2CC')
fill_sub=PatternFill('solid',fgColor='D9E1F2'); fill_tot=PatternFill('solid',fgColor='EDEDED')
STF={'MEC':'DDEBF7','ELE':'FCE4D6','PLAT':'E2EFDA','CDP':'FFF2CC','NORM':'EDE2F6'}
thin=Side(style='thin',color='BFBFBF'); box=Border(left=thin,right=thin,top=thin,bottom=thin)
center=Alignment(horizontal='center',vertical='center',wrap_text=True)
NM=9; START=dt.date(2026,10,1)
def months():
    y,m=START.year,START.month
    for i in range(NM):
        yield dt.date(y+(m-1+i)//12,(m-1+i)%12+1,1)
MES=list(months())
def hdr(ws,r,c,v,fill=fill_h,font=f_h):
    x=ws.cell(r,c,v); x.font=font; x.fill=fill; x.alignment=center; x.border=box; return x
def setw(ws,widths):
    for k,v in widths.items(): ws.column_dimensions[k].width=v

# ---------------- PARAMETROS ----------------
P=wb.active; P.title='PARAMETROS'
P['A1']='PARÂMETROS DE CAPACIDADE'; P['A1'].font=f_t
P['A2']='Células em azul com fundo amarelo são entradas. O resto é calculado. Origem dos valores: aba CAPACIDADE (2) da carteira FGI 8.5-01.02.'; P['A2'].font=Font(name=F,italic=True,size=9)
hdr(P,4,1,'Item'); hdr(P,4,2,'Unid.')
for i in range(NM):
    c=3+i
    x=P.cell(4,c, MES[0] if i==0 else f'=EDATE({L(c-1)}4,1)')
    x.number_format='mmm/yy'; x.font=f_in if i==0 else f_h; x.fill=fill_in if i==0 else fill_h; x.alignment=center; x.border=box
P['C4'].comment=Comment('Mês inicial do horizonte do PMP (entrada). Os demais meses seguem automaticamente.','PMP')
#               out/26 nov  dez  jan27 fev mar abr mai jun
param=[
 ('Horas por dia','h',      [9]*9),
 ('Eficiência','%',         [0.78,0.78,0.78,0.9,0.9,0.9,0.9,0.9,0.9]),
 ('Dias úteis','dias',      [21,19,15,15,18,22,21,20,22]),
 ('Mecânica','pessoas',     [3,3,3,3,3,3,3,3,3]),
 ('Elétrica (próprios)','pessoas',[2]*9),
 ('Elétrica (terceiros)','pessoas',[9]*9),
 ('Plataforma (giga)','pessoas',[2]*9),
 ('CDP (testes)','pessoas', [3]*9),
 ('Normalização/Insp./Emb.','pessoas',[2]*9),
]
r=5
for name,u,vals in param:
    P.cell(r,1,name).font=f_n; P.cell(r,2,u).font=f_n
    for i,v in enumerate(vals):
        x=P.cell(r,3+i,v); x.font=f_in; x.fill=fill_in; x.border=box
        x.number_format='0%' if u=='%' else '0.0' if u=='pessoas' else '0'
    r+=1
# assumptions flags
for col in range(6,12):  # jan27..jun27 headcount & days are my assumptions
    pass
for c in range(6,12):
    P.cell(7,c).comment=Comment('Premissa (não estava na carteira): jan = 15 dias por férias coletivas como em jan/26; demais = dias úteis do calendário menos feriados nacionais. Ajuste.','PMP')
    for rr in range(8,14):
        P.cell(rr,c).comment=Comment('Premissa: carteira não tinha headcount para 2027 — repeti o quadro de dez/26. Ajuste.','PMP')
P['A15']='Fator capacidade flex (hora extra/turno)'; P['A15'].font=f_n
P['C15']=1.2; P['C15'].font=f_in; P['C15'].fill=fill_in; P['C15'].number_format='0.00'
P['C15'].comment=Comment('Mesmo fator 1,2 usado na aba CAPACIDADE (2) da carteira.','PMP')
hdr(P,17,1,'CAPACIDADE (h/mês)'); hdr(P,17,2,'Centro')
for i in range(NM):
    x=P.cell(17,3+i,f'={L(3+i)}4'); x.number_format='mmm/yy'; x.font=f_h; x.fill=fill_h; x.alignment=center
caps=[('Mecânica','MEC','{c}8'),('Elétrica','ELE','({c}9+{c}10)'),('Plataforma','PLAT','{c}11'),('CDP','CDP','{c}12'),('Normalização','NORM','{c}13')]
CAPROW={}
for k,(name,code,ref) in enumerate(caps):
    rr=18+k; CAPROW[code]=rr
    P.cell(rr,1,name).font=f_n; P.cell(rr,2,code).font=f_b
    for i in range(NM):
        c=L(3+i); x=P.cell(rr,3+i,f'={c}5*{c}6*{c}7*'+ref.format(c=c)); x.number_format='#,##0'; x.font=f_n; x.border=box
P.cell(23,1,'TOTAL').font=f_b
for i in range(NM):
    c=L(3+i); x=P.cell(23,3+i,f'=SUM({c}18:{c}22)'); x.number_format='#,##0'; x.font=f_b; x.fill=fill_tot; x.border=box
setw(P,{'A':38,'B':10,**{L(3+i):11 for i in range(NM)}})
P.freeze_panes='C5'

# ---------------- TEMPOS ----------------
T=wb.create_sheet('TEMPOS')
T['A1']='TEMPOS PADRÃO POR ETAPA (h/painel)'; T['A1'].font=f_t
T['A2']='Fonte: aba TEMPOS da carteira (tabelas MEC, ELE, PLATAFORMA, CDP|EMBALAGEM, NORM/INSP/EMB). Edite aqui para recalcular todo o PMP.'; T['A2'].font=Font(name=F,italic=True,size=9)
STG=['MEC','ELE','PLAT','CDP','NORM']; STK=['MEC','ELE','PLAT','CDP','NORM/INSP/EMB']
for j,h in enumerate(['Tipo','Complexidade','Chave']+STG+['Total']): hdr(T,4,1+j,h)
for i,(tp,cx,key) in enumerate(D['keys']):
    rr=5+i
    T.cell(rr,1,tp).font=f_n; T.cell(rr,2,cx).font=f_n; T.cell(rr,3,f'=A{rr}&B{rr}').font=f_n
    for j,s in enumerate(STK):
        x=T.cell(rr,4+j,D['tempo'][key][s]); x.font=f_in; x.fill=fill_in; x.border=box
    T.cell(rr,9,f'=SUM(D{rr}:H{rr})').font=f_b
TLAST=5+len(D['keys'])-1
setw(T,{'A':10,'B':14,'C':16,'D':8,'E':8,'F':8,'G':8,'H':8,'I':8})

# ---------------- PMP ----------------
M=wb.create_sheet('PMP',0)
M['A1']='PMP — PLANO MESTRE DE PRODUÇÃO (1 linha por painel)'; M['A1'].font=f_t
M['A2']='Azul = entrada. Preencha o MÊS de cada etapa (1º dia do mês) e o % RESTANTE (1 = 100% a fazer; 0 = etapa concluída). Horas vêm de TEMPOS × %.'; M['A2'].font=Font(name=F,italic=True,size=9)
M['A3']='Mês inicial do horizonte:'; M['A3'].font=f_b
M['D3']='=PARAMETROS!C4'; M['D3'].number_format='mmm/yy'; M['D3'].font=f_link
base=['OS','Projeto','Subestação','TAG','Tipo','Complexidade','Chave','Testes','Giga','Origem','Data alvo (entrega)','Doc. liberada (S/N)','Material crítico OK (S/N)']
for j,h in enumerate(base): hdr(M,5,1+j,h); M.merge_cells(start_row=4,start_column=1+j,end_row=4,end_column=1+j) if False else None
SC={}  # stage -> (mes col, pct col, h col)
c=14
for s in STG:
    SC[s]=(c,c+1,c+2)
    M.merge_cells(start_row=4,start_column=c,end_row=4,end_column=c+2)
    hdr(M,4,c,s,fill=PatternFill('solid',fgColor=STF[s]),font=f_b)
    for k,h in enumerate(['Mês','% rest.','Horas']):
        hdr(M,5,c+k,h,fill=PatternFill('solid',fgColor=STF[s]),font=f_b)
    c+=3
calc=['Horas restantes','Mês de saída','Sequência','Prazo vs alvo','Liberação','Reprogramar?']
CC={}
for k,h in enumerate(calc):
    hdr(M,5,c+k,h); CC[h]=c+k
hdr(M,4,c,'VERIFICAÇÕES'); M.merge_cells(start_row=4,start_column=c,end_row=4,end_column=c+len(calc)-1)
LASTCOL=c+len(calc)-1
ROW0=6; MAXR=800
rows=D['rows']
yn=DataValidation(type='list',formula1='"S,N"',allow_blank=True); M.add_data_validation(yn)
for i in range(MAXR-ROW0+1):
    rr=ROW0+i
    d=rows[i] if i<len(rows) else None
    if d:
        vals=[d['OS'],d['proj'],d['se'],d['tag'],d['tipo'],d['comp'],None,d['testes'],d['giga'],d['orig'],
              dt.date.fromisoformat(d['alvo']) if d.get('alvo') else None,d.get('doc'),d.get('mat')]
        for j,v in enumerate(vals):
            if j==6: continue
            x=M.cell(rr,1+j,v); x.font=f_in
            if j==10: x.number_format='dd/mm/yy'
        for s,sk in zip(STG,STK):
            st=d.get(sk)
            if st:
                if st[0]: x=M.cell(rr,SC[s][0],dt.date.fromisoformat(st[0])); x.font=f_in
                x=M.cell(rr,SC[s][1],st[1]); x.font=f_in
    M.cell(rr,7,f'=IF(E{rr}="","",E{rr}&F{rr})').font=f_n
    for k,s in enumerate(STG):
        mc,pc,hc=SC[s]
        M.cell(rr,mc).number_format='mmm/yy'; M.cell(rr,pc).number_format='0%'
        tc=L(4+k)
        x=M.cell(rr,hc,f'=IF({L(pc)}{rr}="",0,IFERROR(INDEX(TEMPOS!${tc}$5:${tc}${TLAST},MATCH($G{rr},TEMPOS!$C$5:$C${TLAST},0))*{L(pc)}{rr},0))')
        x.number_format='0;-0;-'; x.font=f_n
    hs='+'.join(f'{L(SC[s][2])}{rr}' for s in STG)
    M.cell(rr,CC['Horas restantes'],f'={hs}').number_format='0;-0;-'
    ms=','.join(f'{L(SC[s][0])}{rr}' for s in STG)
    x=M.cell(rr,CC['Mês de saída'],f'=IF(COUNT({ms})=0,"",MAX({ms}))'); x.number_format='mmm/yy'
    m={s:f'{L(SC[s][0])}{rr}' for s in STG}
    pairs=[('MEC','ELE'),('MEC','PLAT'),('MEC','CDP'),('MEC','NORM'),('ELE','PLAT'),('ELE','CDP'),('ELE','NORM'),('PLAT','NORM'),('CDP','NORM')]
    conds=','.join(f'AND({m[a]}<>"",{m[b]}<>"",{m[a]}>{m[b]})' for a,b in pairs)
    M.cell(rr,CC['Sequência'],f'=IF($A{rr}="","",IF(OR({conds}),"FORA DE ORDEM","OK"))')
    sa=L(CC['Mês de saída'])
    M.cell(rr,CC['Prazo vs alvo'],f'=IF(OR($K{rr}="",{sa}{rr}=""),"",IF({sa}{rr}>DATE(YEAR($K{rr}),MONTH($K{rr}),1),"ATRASA","OK"))')
    mec=m['MEC']
    M.cell(rr,CC['Liberação'],f'=IF($A{rr}="","",IF(AND($L{rr}="S",$M{rr}="S"),"LIBERADO",IF({L(SC['MEC'][2])}{rr}=0,"-",IF(AND({mec}<>"",{mec}<=EDATE(PARAMETROS!$C$4,1)),"RISCO","PENDENTE"))))')
    past=','.join(f'AND({L(SC[s][0])}{rr}<>"",{L(SC[s][0])}{rr}<PARAMETROS!$C$4,{L(SC[s][2])}{rr}>0)' for s in STG)
    nom=','.join(f'AND({L(SC[s][0])}{rr}="",{L(SC[s][2])}{rr}>0)' for s in STG)
    M.cell(rr,CC['Reprogramar?'],f'=IF($A{rr}="","",IF(OR({past}),"MÊS PASSADO",IF(OR({nom}),"SEM MÊS","")))')
    for cc in range(1,LASTCOL+1):
        x=M.cell(rr,cc); x.border=box
        if x.font!=f_in: x.font=x.font if x.font.color and x.font.color.rgb=='FF0000FF' else f_n
    for cc in (12,13): yn.add(M.cell(rr,cc))
NR=len(rows)
red=PatternFill('solid',fgColor='F8CBAD'); amb=PatternFill('solid',fgColor='FFE699'); grn=PatternFill('solid',fgColor='C6EFCE')
rng=f'{L(CC["Sequência"])}{ROW0}:{L(LASTCOL)}{MAXR}'
M.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=['"FORA DE ORDEM"'],fill=red))
M.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=['"ATRASA"'],fill=red))
M.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=['"RISCO"'],fill=red))
M.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=['"MÊS PASSADO"'],fill=amb))
M.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=['"SEM MÊS"'],fill=amb))
M.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=['"PENDENTE"'],fill=amb))
M.conditional_formatting.add(rng,CellIsRule(operator='equal',formula=['"LIBERADO"'],fill=grn))
M.freeze_panes=M.cell(ROW0,5)
M.auto_filter.ref=f'A5:{L(LASTCOL)}{ROW0+NR-1}'
setw(M,{'A':7,'B':34,'C':20,'D':14,'E':7,'F':10,'G':12,'H':12,'I':6,'J':7,'K':10,'L':8,'M':9})
for s in STG:
    for k,w in zip(SC[s],(8,6,6)): M.column_dimensions[L(k)].width=w
for h,w in zip(calc,(9,9,14,10,10,13)): M.column_dimensions[L(CC[h])].width=w
M.column_dimensions['G'].hidden=True
M.row_dimensions[5].height=42

# ---------------- CARGA x CAPACIDADE ----------------
C=wb.create_sheet('CARGA x CAPACIDADE',1)
C['A1']='CARGA × CAPACIDADE POR CENTRO (horas)'; C['A1'].font=f_t
C['A2']='Carga = soma das horas do PMP no mês. "Backlog acumulado" = horas que não couberam e empurram para o mês seguinte (inclui o que está em mês passado). Verde ≤85% · Amarelo ≤100% · Laranja ≤ flex · Vermelho > flex.'; C['A2'].font=Font(name=F,italic=True,size=9)
hdr(C,4,1,'Centro'); hdr(C,4,2,'Linha'); hdr(C,4,3,'Sem mês'); hdr(C,4,4,'Mês passado')
for i in range(NM):
    x=C.cell(4,5+i,f'=PARAMETROS!{L(3+i)}4'); x.number_format='mmm/yy'; x.font=f_h; x.fill=fill_h; x.alignment=center
hdr(C,4,5+NM,'Total horizonte')
LC=5+NM
names={'MEC':'Mecânica','ELE':'Elétrica','PLAT':'Plataforma','CDP':'CDP','NORM':'Normalização'}
r=5; UT=[]; BLK={}
R1,R2=ROW0,MAXR
for s in STG+['TOTAL']:
    BLK[s]=r
    labels=['Capacidade','Capacidade flex','Carga','Utilização','Saldo (cap − carga)','Backlog acumulado']
    for k,lab in enumerate(labels):
        rr=r+k
        C.cell(rr,1,names.get(s,'TOTAL') if k==0 else None).font=f_b
        C.cell(rr,2,lab).font=f_b if k in (2,3) else f_n
        for i in range(NM):
            col=5+i; cl=L(col)
            if s!='TOTAL':
                mc,_,hc=SC[s]
                fm={0:f'=PARAMETROS!{L(3+i)}{CAPROW[s]}',
                    1:f'={cl}{r}*PARAMETROS!$C$15',
                    2:f'=SUMIFS(PMP!${L(hc)}${R1}:${L(hc)}${R2},PMP!${L(mc)}${R1}:${L(mc)}${R2},{cl}$4)'}
            else:
                fm={k2:'='+'+'.join(f'{cl}{BLK[x]+k2}' for x in STG) for k2 in (0,1,2)}
            fm[3]=f'=IFERROR({cl}{r+2}/{cl}{r},0)'
            fm[4]=f'={cl}{r}-{cl}{r+2}'
            fm[5]=(f'=MAX(0,$D{r+2}+{cl}{r+2}-{cl}{r})' if i==0 else f'=MAX(0,{L(col-1)}{rr}+{cl}{r+2}-{cl}{r})')
            x=C.cell(rr,col,fm[k]); x.number_format='0%' if k==3 else '#,##0;(#,##0);-'; x.font=f_link if k==0 else f_n; x.border=box
        if k==2:
            if s!='TOTAL':
                mc,_,hc=SC[s]
                C.cell(rr,3,f'=SUMPRODUCT((PMP!${L(mc)}${R1}:${L(mc)}${R2}="")*PMP!${L(hc)}${R1}:${L(hc)}${R2})')
                C.cell(rr,4,f'=SUMIFS(PMP!${L(hc)}${R1}:${L(hc)}${R2},PMP!${L(mc)}${R1}:${L(mc)}${R2},"<"&PARAMETROS!$C$4)')
            else:
                C.cell(rr,3,'='+'+'.join(f'C{BLK[x]+2}' for x in STG)); C.cell(rr,4,'='+'+'.join(f'D{BLK[x]+2}' for x in STG))
            for cc in (3,4): C.cell(rr,cc).number_format='#,##0;(#,##0);-'; C.cell(rr,cc).border=box; C.cell(rr,cc).font=f_n
        if k in (0,1,2,4):
            x=C.cell(rr,LC,f'=SUM({L(5)}{rr}:{L(4+NM)}{rr})'); x.number_format='#,##0;(#,##0);-'; x.font=f_b; x.border=box
        if k==3:
            x=C.cell(rr,LC,f'=IFERROR({L(LC)}{r+2}/{L(LC)}{r},0)'); x.number_format='0%'; x.font=f_b; x.border=box
            UT.append(rr)
        if k==2:
            for cc in range(1,LC+1): C.cell(rr,cc).fill=fill_sub
    r+=7
for rr in UT:
    ref=f'E{rr}:{L(LC)}{rr}'
    C.conditional_formatting.add(ref,FormulaRule(formula=[f'E{rr}>PARAMETROS!$C$15'],fill=PatternFill('solid',fgColor='FF7C80'),stopIfTrue=True))
    C.conditional_formatting.add(ref,FormulaRule(formula=[f'E{rr}>1'],fill=PatternFill('solid',fgColor='F4B183'),stopIfTrue=True))
    C.conditional_formatting.add(ref,FormulaRule(formula=[f'E{rr}>0.85'],fill=PatternFill('solid',fgColor='FFE699'),stopIfTrue=True))
    C.conditional_formatting.add(ref,FormulaRule(formula=[f'E{rr}>0'],fill=PatternFill('solid',fgColor='C6EFCE'),stopIfTrue=True))
setw(C,{'A':14,'B':20,'C':10,'D':11,**{L(5+i):10 for i in range(NM)},L(LC):13})
C.freeze_panes='E5'
# charts
cr=r+1
for n,s in enumerate(STG+['TOTAL']):
    b=BLK[s]
    bar=BarChart(); bar.type='col'; bar.title=f'{names.get(s,"TOTAL")} — carga × capacidade'; bar.y_axis.title='h'
    bar.add_data(Reference(C,min_col=4,max_col=4+NM,min_row=b+2),titles_from_data=False,from_rows=True)
    bar.series[0].tx=None
    from openpyxl.chart.series import SeriesLabel
    bar.series[0].tx=SeriesLabel(v='Carga (inclui mês passado)')
    bar.set_categories(Reference(C,min_col=4,max_col=4+NM,min_row=4))
    ln=LineChart()
    for k,lab in ((0,'Capacidade'),(1,'Capacidade flex')):
        ln.add_data(Reference(C,min_col=4,max_col=4+NM,min_row=b+k),titles_from_data=False,from_rows=True)
        ln.series[-1].tx=SeriesLabel(v=lab)
    bar+=ln; bar.height=7; bar.width=16; bar.legend.position='b'
    C.add_chart(bar,f'{"A" if n%2==0 else "H"}{cr+(n//2)*15}')

# ---------------- PMP MENSAL ----------------
S=wb.create_sheet('PMP MENSAL',2)
S['A1']='PMP MENSAL — painéis com saída (última etapa) por projeto e mês'; S['A1'].font=f_t
S['A2']='Saída = mês da última etapa planejada do painel (coluna "Mês de saída" do PMP). Bloco de baixo: horas restantes por projeto/mês (todas as etapas).'; S['A2'].font=Font(name=F,italic=True,size=9)
projs=sorted({x['proj'] for x in rows})
def block(top,title,kind):
    hdr(S,top,1,title); hdr(S,top,2,'Mês passado')
    for i in range(NM):
        x=S.cell(top,3+i,f'=PARAMETROS!{L(3+i)}4'); x.number_format='mmm/yy'; x.font=f_h; x.fill=fill_h; x.alignment=center
    hdr(S,top,3+NM,'Total')
    sa=L(CC['Mês de saída'])
    for j,p in enumerate(projs):
        rr=top+1+j
        S.cell(rr,1,p).font=f_n; S.cell(rr,1).border=box
        for i in range(-1,NM):
            col=3+i; cl=L(col)
            if kind=='cnt':
                f=(f'=COUNTIFS(PMP!$B${R1}:$B${R2},$A{rr},PMP!${sa}${R1}:${sa}${R2},"<"&PARAMETROS!$C$4)' if i<0 else
                   f'=COUNTIFS(PMP!$B${R1}:$B${R2},$A{rr},PMP!${sa}${R1}:${sa}${R2},{cl}${top})')
            else:
                parts=[]
                for s in STG:
                    mc,_,hc=SC[s]
                    crit=f'"<"&PARAMETROS!$C$4' if i<0 else f'{cl}${top}'
                    parts.append(f'SUMIFS(PMP!${L(hc)}${R1}:${L(hc)}${R2},PMP!$B${R1}:$B${R2},$A{rr},PMP!${L(mc)}${R1}:${L(mc)}${R2},{crit})')
                f='='+'+'.join(parts)
            x=S.cell(rr,col,f); x.number_format='#,##0;(#,##0);-'; x.font=f_n; x.border=box
        x=S.cell(rr,3+NM,f'=SUM(B{rr}:{L(2+NM)}{rr})'); x.font=f_b; x.number_format='#,##0;(#,##0);-'; x.border=box
    tr=top+1+len(projs)
    S.cell(tr,1,'TOTAL').font=f_b
    for col in range(2,4+NM):
        x=S.cell(tr,col,f'=SUM({L(col)}{top+1}:{L(col)}{tr-1})'); x.font=f_b; x.fill=fill_tot; x.number_format='#,##0;(#,##0);-'; x.border=box
    return tr
t1=block(4,'Projeto — PAINÉIS (saída)','cnt')
t2=block(t1+3,'Projeto — HORAS restantes','h')
setw(S,{'A':52,'B':11,**{L(3+i):9 for i in range(NM)},L(3+NM):9})
S.freeze_panes='B5'

# ---------------- LEIAME ----------------
R=wb.create_sheet('LEIAME',0)
txt=[('COMO USAR ESTE PMP',f_t),
('',f_n),
('1. PARAMETROS — ajuste headcount, dias úteis e eficiência de cada mês. Isso gera a capacidade em horas por centro.',f_n),
('2. TEMPOS — horas padrão por Tipo+Complexidade e por etapa (copiado da aba TEMPOS da carteira).',f_n),
('3. PMP — 1 linha por painel. Você só mexe nas colunas azuis: o MÊS de cada etapa e o % RESTANTE.',f_n),
('     • % restante = 1 (100%) quando a etapa não começou; 0,5 se está na metade; 0 quando concluiu (a hora sai da carga).',f_n),
('     • Data alvo, Doc. liberada e Material crítico OK alimentam os alertas de prazo e liberação. Liberação = RISCO quando a MEC começa em até 2 meses e Doc/Material não estão OK.',f_n),
('4. CARGA x CAPACIDADE — mostra mês a mês se o plano cabe. Se a utilização passa de 100% (ou do flex), o plano é irreal: mova painéis no PMP até ficar verde/amarelo.',f_n),
('5. PMP MENSAL — o compromisso: quantos painéis de cada projeto saem em cada mês. É esta visão que vai para Comercial/PM.',f_n),
('',f_n),
('ROTINA SUGERIDA',f_b),
('• Semanal (PCP): atualizar % restante com base no Kanban de andamento; resolver alertas "MÊS PASSADO" e "SEM MÊS".',f_n),
('• Mensal (reunião S&OP): congelar o mês seguinte (não mexe mais) e nivelar M+2 e M+3 contra a capacidade.',f_n),
('',f_n),
('REGRAS DE SEQUÊNCIA VERIFICADAS',f_b),
('MEC ≤ ELE ≤ CDP ≤ NORM · ELE ≤ PLAT ≤ NORM (PLAT pode ser antes ou depois do CDP — padrão observado na carteira).',f_n),
('',f_n),
('ORIGEM DOS DADOS',f_b),
('Carteira FGI 8.5-01.02 (rev. 08): só linhas com "PROJETO ATIVO PLANEJAMENTO" = SIM e ORIGEM ≠ CANCELADO, agrupadas por OS.',f_n),
('O % restante inicial = Horas/Mês da carteira ÷ hora padrão (reproduz exatamente as horas da carteira).',f_n),
('Data alvo veio do "Cronograma SET-OUT" só quando a Meta era EMBALADO; Doc/Material também de lá. O resto está em branco — preencher.',f_n),
('Células com comentário em PARAMETROS (2027) são premissas minhas — a carteira não tinha headcount para 2027.',f_n),
('',f_n),
('Novo painel: cole na próxima linha vazia do PMP (fórmulas já estão até a linha 800).',f_n),
]
for i,(t,fn) in enumerate(txt): R.cell(1+i,1,t).font=fn
R.column_dimensions['A'].width=140
wb.active=0
wb.save(OUT)
print('ok',NR)
